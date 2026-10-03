"""Lưu JSON và tạo bảng experiments.xlsx.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    out=Path(results_dir); out.mkdir(parents=True,exist_ok=True); p=out/f"{result['cfg']['exp_id']}.json"
    with p.open("w",encoding="utf-8") as f: json.dump({k:result[k] for k in ("cfg","history","summary")},f,indent=2,ensure_ascii=False)
    return str(p)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    root=Path(results_dir)
    if not root.exists(): return []
    values=[]
    for p in sorted(root.glob("*.json")):
        with p.open(encoding="utf-8") as f: values.append(json.load(f))
    return sorted(values,key=lambda r:r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    row={**result["cfg"],**result["summary"]}; row["hidden"]="-".join(map(str,result["cfg"]["hidden"]))
    row["eval_acc"]=eval_scores.get("accuracy") if eval_scores else None
    row["eval_macro_f1"]=eval_scores.get("macro_f1") if eval_scores else None
    row["figure_file"]=f"figures/{result['cfg']['exp_id']}.png"; row["notes"]=notes
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    import openpyxl
    wb=openpyxl.load_workbook(template_path); ws=wb["Experiments"]; headers=[c.value for c in ws[1]]
    formulas={"step0_gap_vs_lnC","gap_val_minus_train","delta_val_f1_vs_base","beyond_noise"}
    for i,row in enumerate(rows,2):
        for j,key in enumerate(headers,1):
            if key in row and key not in formulas:
                v=row[key]
                if isinstance(v,(list,tuple,dict)): v=json.dumps(v,ensure_ascii=False)
                ws.cell(i,j).value=v
    out=Path(out_path); out.parent.mkdir(parents=True,exist_ok=True); wb.save(out)
